#!/usr/bin/env python3
"""
Build the two KoboToolbox XLSForms for the facility project:

  forms/Capture_All_Facilities.xlsx
      Section 1  Access  - collector name + private access code (password)
      Section 2  Location - LGA, Ward, Community (cascading) + GPS
      Section 3  Facility - name, type, picture of the facility
      Section 4  Review

  forms/Facility_Report_Geofenced.xlsx
      Section 1  Access   - collector name + private access code (password)
      Section 2  Facility - choose the registered facility (from facilities.csv)
      Section 3  On-site check - GPS must be within 100 m of the facility
      Section 4  Report   - operational status, condition, notes, photo
      Section 5  Review

Every generated XLSForm is converted with pyxform (the same converter
KoboToolbox runs on deployment) so you know it will deploy cleanly.

Usage:
    python scripts/build_forms.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    CHOICES_COLS,
    CONDITION_RATINGS,
    FACILITY_TYPES,
    FORMS_DIR,
    OPERATIONAL_STATUS,
    SURVEY_COLS,
    check_media_csvs,
    crow,
    end_group,
    group,
    load_locations,
    srow,
)

# --------------------------------------------------------------------------- #
# Shared expressions
# --------------------------------------------------------------------------- #
ACCESS_CODE_CHECK = (
    "string(.) = pulldata('enumerators', 'password', 'name', ${collector})"
)
ACCESS_OK = (
    "if(string-length(${access_code}) > 0 and "
    "${access_code} = pulldata('enumerators', 'password', 'name', ${collector}), 1, 0)"
)
ACCESS_GRANTED = "number(${access_ok}) = 1"
GPS_PARAMS = "capture-accuracy=10 warning-accuracy=50"
GEOFENCE_RADIUS_M = 100  # metres - appears in the constraint below (Form 2)

GEOFENCE_CONSTRAINT = (
    "if(string-length(${fac_latitude}) > 0, "
    "distance(${current_gps}, ${fac_point}) <= "
    f"{GEOFENCE_RADIUS_M}, true())"
)
GEOFENCE_MESSAGE = (
    f"You must be within {GEOFENCE_RADIUS_M} metres of the selected facility "
    "to submit this report. Please move closer and get your GPS position again."
)


def access_section():
    """Name + password gate - identical in both forms."""
    return [
        group("grp_access", "Section 1: Access - collector name and password"),
        srow(
            type="select_one_from_file enumerators.csv",
            name="collector",
            label="Name of collector (who is collecting this report?)",
            required="yes",
            hint="Select your name from the list.",
        ),
        srow(
            type="text",
            name="access_code",
            label="Access code (password)",
            relevant="${collector} != ''",
            required="yes",
            constraint=ACCESS_CODE_CHECK,
            constraint_message=(
                "Wrong access code for the selected name. "
                "Re-type it carefully or contact your supervisor."
            ),
            appearance="masked",
            hint=(
                "Enter the private code issued to you. "
                "In KoboCollect the characters are hidden while typing."
            ),
        ),
        srow(type="calculate", name="access_ok", calculation=ACCESS_OK),
        srow(
            type="note",
            name="access_granted",
            label="✅ Access granted. Continue to the next section.",
            relevant=ACCESS_GRANTED,
        ),
        end_group(),
    ]


def meta_section():
    return [
        srow(type="start", name="start"),
        srow(type="end", name="end"),
        srow(type="deviceid", name="deviceid"),
        srow(type="username", name="username"),
    ]


# --------------------------------------------------------------------------- #
# Form 1 - Capture all facilities
# --------------------------------------------------------------------------- #
def build_capture_form() -> Workbook:
    locs = load_locations()

    survey = access_section()

    survey += [
        group("grp_location", "Section 2: Location of the facility", ACCESS_GRANTED),
        srow(
            type="select_one lga_list",
            name="lga_sel",
            label="LGA",
            required="yes",
        ),
        srow(
            type="select_one ward_list",
            name="ward_sel",
            label="Ward",
            required="yes",
            choice_filter="lga = ${lga_sel}",
        ),
        srow(
            type="select_one community_list",
            name="community_sel",
            label="Community",
            required="yes",
            choice_filter="lga = ${lga_sel} and ward = ${ward_sel}",
        ),
        srow(
            type="calculate",
            name="lga_label",
            calculation="instance('lga_list')/root/item[name=${lga_sel}]/label",
        ),
        srow(
            type="calculate",
            name="ward_label",
            calculation="instance('ward_list')/root/item[name=${ward_sel}]/label",
        ),
        srow(
            type="calculate",
            name="community_label",
            calculation="instance('community_list')/root/item[name=${community_sel}]/label",
        ),
        srow(
            type="geopoint",
            name="facility_gps",
            label="GPS coordinates of the facility",
            required="yes",
            parameters=GPS_PARAMS,
            hint=(
                "Stand at the facility (or its gate) and wait for the position "
                "to be recorded. This proves the facility was captured on site."
            ),
        ),
        end_group(),
        group("grp_facility", "Section 3: Facility details", ACCESS_GRANTED),
        srow(
            type="text",
            name="facility_name",
            label="Facility name (as written on the signboard)",
            required="yes",
        ),
        srow(
            type="select_one facility_type_list",
            name="facility_type",
            label="Type of facility",
            required="yes",
        ),
        srow(
            type="calculate",
            name="facility_type_label",
            calculation=(
                "instance('facility_type_list')/root/item[name=${facility_type}]/label"
            ),
        ),
        srow(
            type="image",
            name="facility_photo",
            label="Picture of the facility",
            required="yes",
            parameters="max-pixels=1600",
            hint="Take a clear photo of the building or its signboard.",
        ),
        end_group(),
        group("grp_review", "Section 4: Review and submit", ACCESS_GRANTED),
        srow(
            type="note",
            name="review_note",
            label=(
                "**Please review your answers**\n\n"
                "Collector: ${collector}\n\n"
                "Facility: ${facility_name}\n\n"
                "Type: ${facility_type_label}\n\n"
                "LGA: ${lga_label}"
                " | Ward: ${ward_label}"
                " | Community: ${community_label}\n\n"
                "GPS: ${facility_gps}\n\n"
                "When everything is correct, submit the form."
            ),
        ),
        end_group(),
    ]
    survey += meta_section()

    choices = []
    for lga_id, lga_label in locs.lgas:
        choices.append(crow("lga_list", lga_id, lga_label))
    for ward_id, ward_label, lga_id in locs.wards:
        choices.append(crow("ward_list", ward_id, ward_label, lga=lga_id))
    for comm_id, comm_label, lga_id, ward_id in locs.communities:
        choices.append(
            crow("community_list", comm_id, comm_label, lga=lga_id, ward=ward_id)
        )
    for name, label in FACILITY_TYPES:
        choices.append(crow("facility_type_list", name, label))

    return _assemble(
        title="Capture All Facilities - On-site Registration",
        form_id="capture-all-facilities",
        version="1.0.0",
        survey=survey,
        choices=choices,
    )


# --------------------------------------------------------------------------- #
# Form 2 - On-site report, geofenced to the facility GPS
# --------------------------------------------------------------------------- #
def build_report_form() -> Workbook:
    survey = access_section()
    survey += [
        group(
            "grp_facility",
            "Section 2: Which facility are you reporting on?",
            ACCESS_GRANTED,
        ),
        srow(
            type="select_one_from_file facilities.csv",
            name="facility_sel",
            label="Facility",
            required="yes",
            appearance="map",
            hint="Select the facility from the list, or pick its marker on the map.",
        ),
        srow(
            type="calculate",
            name="fac_name",
            calculation="pulldata('facilities', 'label', 'name', ${facility_sel})",
        ),
        srow(
            type="calculate",
            name="fac_lga",
            calculation="pulldata('facilities', 'lga', 'name', ${facility_sel})",
        ),
        srow(
            type="calculate",
            name="fac_ward",
            calculation="pulldata('facilities', 'ward', 'name', ${facility_sel})",
        ),
        srow(
            type="calculate",
            name="fac_community",
            calculation="pulldata('facilities', 'community', 'name', ${facility_sel})",
        ),
        srow(
            type="calculate",
            name="fac_type",
            calculation="pulldata('facilities', 'facility_type', 'name', ${facility_sel})",
        ),
        srow(
            type="calculate",
            name="fac_latitude",
            calculation="pulldata('facilities', 'latitude', 'name', ${facility_sel})",
        ),
        srow(
            type="calculate",
            name="fac_longitude",
            calculation="pulldata('facilities', 'longitude', 'name', ${facility_sel})",
        ),
        srow(
            type="calculate",
            name="fac_point",
            calculation=(
                "if(string-length(${fac_latitude}) > 0, "
                "concat(${fac_latitude}, ' ', ${fac_longitude}, ' 0 100'), '')"
            ),
        ),
        srow(
            type="note",
            name="facility_info",
            label=(
                "**Facility details (from the register)**\n\n"
                "${fac_name}\n\n"
                "Type: ${fac_type}\n\n"
                "Location: ${fac_community}, ${fac_ward}, ${fac_lga}"
            ),
            relevant="${facility_sel} != ''",
        ),
        end_group(),
        group(
            "grp_geofence",
            "Section 3: On-site GPS check - you must be at the facility",
            ACCESS_GRANTED,
        ),
        srow(
            type="geopoint",
            name="current_gps",
            label="Your current GPS position",
            relevant="${facility_sel} != ''",
            required="yes",
            constraint=GEOFENCE_CONSTRAINT,
            constraint_message=GEOFENCE_MESSAGE,
            parameters=GPS_PARAMS,
            hint=(
                f"Get your position while standing at the facility. "
                f"The form only accepts it within {GEOFENCE_RADIUS_M} m "
                "of the registered facility coordinates."
            ),
        ),
        srow(
            type="calculate",
            name="distance_m",
            calculation=(
                "if(string-length(${fac_latitude}) > 0, "
                "if(${current_gps} != '', "
                "round(distance(${current_gps}, ${fac_point})), ''), '')"
            ),
        ),
        srow(
            type="note",
            name="distance_note",
            label=(
                f"📍 Distance from the facility: **${{distance_m}} metres** "
                f"(allowed radius: {GEOFENCE_RADIUS_M} m)."
            ),
            relevant=(
                "${facility_sel} != '' and string-length(${fac_latitude}) > 0 "
                "and ${current_gps} != ''"
            ),
        ),
        end_group(),
        group(
            "grp_report",
            "Section 4: Facility report",
            "number(${distance_m}) <= 100 and string-length(${fac_latitude}) > 0",
        ),
        srow(
            type="select_one operational_status_list",
            name="operational_status",
            label="Current operational status of the facility",
            required="yes",
        ),
        srow(
            type="calculate",
            name="operational_status_label",
            calculation=(
                "instance('operational_status_list')/root/item"
                "[name=${operational_status}]/label"
            ),
        ),
        srow(
            type="select_one condition_list",
            name="building_condition",
            label="Overall condition of the facility",
            required="yes",
        ),
        srow(
            type="calculate",
            name="condition_label",
            calculation=(
                "instance('condition_list')/root/item"
                "[name=${building_condition}]/label"
            ),
        ),
        srow(
            type="text",
            name="report_notes",
            label="Observations / comments",
            hint="Add anything the supervisor should know about this facility.",
        ),
        srow(
            type="image",
            name="visit_photo",
            label="Additional photo (optional)",
            parameters="max-pixels=1600",
            hint="e.g. a second angle, equipment, or a problem you observed.",
        ),
        end_group(),
        group("grp_review", "Section 5: Review and submit", ACCESS_GRANTED),
        srow(
            type="note",
            name="review_note",
            label=(
                "**Please review your answers**\n\n"
                "Collector: ${collector}\n\n"
                "Facility: ${fac_name} (${fac_type})\n\n"
                "Status: ${operational_status_label}"
                " | Condition: ${condition_label}\n\n"
                f"Distance from facility at GPS capture: ${{distance_m}} m\n\n"
                "When everything is correct, submit the form."
            ),
        ),
        end_group(),
    ]
    survey += meta_section()

    choices = []
    for name, label in OPERATIONAL_STATUS:
        choices.append(crow("operational_status_list", name, label))
    for name, label in CONDITION_RATINGS:
        choices.append(crow("condition_list", name, label))

    return _assemble(
        title="Facility Report - On-site Only (GPS geofenced)",
        form_id="facility-report-geofenced",
        version="1.0.0",
        survey=survey,
        choices=choices,
    )


# --------------------------------------------------------------------------- #
# Assembly + validation
# --------------------------------------------------------------------------- #
def _assemble(title, form_id, version, survey, choices) -> Workbook:
    wb = Workbook()

    ws = wb.active
    ws.title = "survey"
    ws.append(SURVEY_COLS)
    for r in survey:
        ws.append(r)

    ch = wb.create_sheet("choices")
    ch.append(CHOICES_COLS)
    for r in choices:
        ch.append(r)

    st = wb.create_sheet("settings")
    st.append(["form_title", "form_id", "version"])
    st.append([title, form_id, version])

    return wb


def validate(path: Path) -> None:
    """Run pyxform - the converter KoboToolbox itself uses on deployment."""
    from pyxform.xls2xform import convert

    warnings: list = []
    convert(str(path), warnings=warnings)
    status = "OK" if not warnings else "OK (with warnings)"
    print(f"  [pyxform {status}] {path.relative_to(path.parents[1])}")
    for w in warnings:
        print(f"      warning: {w}")


def main() -> None:
    print("Checking external data files ...")
    check_media_csvs()
    load_locations()  # validates data/locations.csv

    FORMS_DIR.mkdir(parents=True, exist_ok=True)

    capture = FORMS_DIR / "Capture_All_Facilities.xlsx"
    report = FORMS_DIR / "Facility_Report_Geofenced.xlsx"

    build_capture_form().save(capture)
    build_report_form().save(report)

    print("Validating with pyxform ...")
    validate(capture)
    validate(report)
    print("Done. Upload the two files in forms/ to KoboToolbox (New project -> XLSForm).")


if __name__ == "__main__":
    main()
